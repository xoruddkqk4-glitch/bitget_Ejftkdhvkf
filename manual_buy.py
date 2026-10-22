# -*- coding: utf-8 -*-
# =========================================================
#  비트겟 떨사오팔 봇 — 수동 매수 전용 스크립트 (manual_buy.py)
# =========================================================
#  ※ 24시간 상시 구동 봇(bitget_Ejftkdhvkf.py)을 멈추지 않고,
#     수동으로 보유금을 보충했거나 추가 매수를 진행하고자 할 때
#     매수 기준을 즉시 확인하고 1회 매수를 안전하게 실행합니다.
#
#  ※ 매수 조건, 투자금 계산, 비트겟 선물 주문, 실체결가 회수,
#     엑셀 매매기록, JSON 상태 등록, 텔레그램/구글시트 알림까지
#     bitget_Ejftkdhvkf.py 와 100% 동일하게 동작합니다.
#
#  ※ 매수 완료 후 엑셀 및 JSON 상태가 갱신되면,
#     24시간 돌아가고 있는 기존 봇이 다음 주기(최대 5분 이내)에
#     신규 차수를 자동으로 인식하여 목표가(+2.5%) 및 2단계 추세추종 매도를 감시합니다.
#
#  [ 실행 방법 ]
#   1) 매수 기준(시가 대비 -1.5% 하락) 점검 후 매수 실행 (기본)
#      python manual_buy.py
#
#   2) 주문 없이 현재 시세, 매수 기준가, 가용 USDT 잔고만 확인
#      python manual_buy.py --check-only
#
#   3) 매수 조건 미달이어도 지금 즉시 1차수 강제 매수
#      python manual_buy.py --force
# =========================================================
import os
import sys
import time
import json
import argparse
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
import pandas as pd
import ccxt

# --- 경로 설정 및 LS공통모듈 탐색 ---
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
LS_MODULE_DIR = os.path.join(BASE_DIR, "LS공통모듈")
if os.path.exists(LS_MODULE_DIR) and LS_MODULE_DIR not in sys.path:
    sys.path.append(LS_MODULE_DIR)

# --- 텔레그램 및 LS 공통 모듈 임포트 ---
import telegram_alert
import LS_Common_new as LsCommon
import LS_API_Helper_US_new as LsiUS

# --- 구글 시트 로거 임포트 ---
from sheets_logger import SheetsLogger
sheet_logger = SheetsLogger(
    webhook_url="https://script.google.com/macros/s/AKfycbws77Zb7aGqs_L5K0yLDf5iDVBXHKO_abf5eL1vlvUTh4197W-dujvmxm-vEBj8Hjdq/exec",
    auth_token="trading_secret_2026",
    strategy_name="떨사오팔",
    exchange="Bitget"
)

# --- 시간 기준: 한국시간(KST) ---
KST = ZoneInfo("Asia/Seoul")

# --- 텔레그램 메시지 접두사 ---
_original_send = telegram_alert.SendMessage
def _prefixed_send(msg):
    return _original_send(f"🤖 [비트겟_떨사오팔|수동매수]\n{msg}")
telegram_alert.SendMessage = _prefixed_send

# --- 비트겟 API 키 로드 ---
try:
    from bitget_keys import BITGET_API_KEY, BITGET_SECRET, BITGET_PASSPHRASE
except ImportError:
    BITGET_API_KEY = "여기에_API_KEY"
    BITGET_SECRET = "여기에_SECRET"
    BITGET_PASSPHRASE = "여기에_PASSPHRASE"
BITGET_API_KEY = BITGET_API_KEY.strip()
BITGET_SECRET = BITGET_SECRET.strip()
BITGET_PASSPHRASE = BITGET_PASSPHRASE.strip()

# --- 심볼 및 주문 파라미터 ---
BITGET_SYMBOL_CANDIDATES = ["SOXL/USDT:USDT", "SOXLUSDT"]
MARGIN_MODE = "cross"
POSITION_SIDE = "long"
USE_LIMIT_ORDER = True
BUY_LIMIT_PRICE_MULTIPLIER = 1.008
ORDER_WAIT_SEC = 2
ORDER_CHECK_COUNT = 5
FILL_FETCH_RETRIES = 5
FILL_FETCH_DELAY = 0.5

# --- 전략 파라미터 (bitget_Ejftkdhvkf.py 와 100% 동일) ---
INITIAL_INVEST_USDT = 150
BUY_DROP_PCT = 0.015             # 시가 대비 -1.5% 하락 시 매수
SELL_PROFIT_PCT = 0.025          # 매수 후 +2.5% 목표가
HIGH_DROP_THRESHOLD = 0.40       # 고점 대비 -40% 하락 시 대폭락 판정
HIGH_DROP_INVEST_MULTIPLIER = 1.5 # 대폭락 시 매수금 1.5배 증액
LOOKBACK_DAYS = 40               # 고점 탐색 기간
DATA_RETRY_COUNT = 5
DATA_RETRY_SLEEP_SEC = 4

# --- 엑셀 및 상태 파일 경로 (서버 우선, 없으면 로컬 자동 전환) ---
SERVER_FILENAME = "/var/autobot/bitget_Ejftkdhvkf/비트겟떨사오팔_매매기록.xlsx"
if os.path.exists(SERVER_FILENAME) or os.path.exists(os.path.dirname(SERVER_FILENAME)):
    FILENAME = SERVER_FILENAME
else:
    FILENAME = os.path.join(BASE_DIR, "비트겟떨사오팔_매매기록.xlsx")

ROUND_STATE_FILE = os.path.join(os.path.dirname(FILENAME), "bitget_round_state.json") if os.path.dirname(FILENAME) else "./bitget_round_state.json"
INVEST_SHEET_NAME = "투자금"

# =========================
# ✅ 비트겟 거래소 인프라 함수
# =========================
def validate_api_strings():
    for name, value in {"API_KEY": BITGET_API_KEY, "SECRET": BITGET_SECRET, "PASSPHRASE": BITGET_PASSPHRASE}.items():
        if not value or "여기에" in value:
            raise RuntimeError(f"비트겟 {name} 가 설정되지 않았습니다. bitget_keys.py 를 확인하세요.")

def create_exchange():
    validate_api_strings()
    exchange = ccxt.bitget({
        "apiKey": BITGET_API_KEY,
        "secret": BITGET_SECRET,
        "password": BITGET_PASSPHRASE,
        "enableRateLimit": True,
        "options": {"defaultType": "swap"},
    })
    exchange.load_markets()
    return exchange

def resolve_symbol(exchange):
    for s in BITGET_SYMBOL_CANDIDATES:
        if s in exchange.markets:
            return s
    raise RuntimeError(f"Bitget에서 SOXL 선물 심볼을 찾지 못했습니다. 후보: {BITGET_SYMBOL_CANDIDATES}")

def get_current_futures_price(exchange, symbol) -> float:
    ticker = exchange.fetch_ticker(symbol)
    last = ticker.get("last")
    if last is None:
        raise RuntimeError("비트겟 현재가 조회 실패")
    return float(last)

def get_available_usdt(exchange) -> float:
    """비트겟 선물(swap) 계좌의 가용 USDT 잔고 조회"""
    try:
        bal = exchange.fetch_balance(params={"type": "swap"})
        free_usdt = None
        if isinstance(bal, dict):
            if "USDT" in bal and isinstance(bal["USDT"], dict):
                free_usdt = bal["USDT"].get("free")
            if free_usdt is None and "free" in bal and isinstance(bal["free"], dict):
                free_usdt = bal["free"].get("USDT")
        if free_usdt is not None:
            return float(free_usdt)
        return 0.0
    except Exception as e:
        print(f"[가용잔고 조회 경고] {e}")
        return -1.0

def usdt_to_amount(exchange, symbol, usdt_amount, price):
    market = exchange.market(symbol)
    raw_amount = float(usdt_amount) / float(price)
    amount = float(exchange.amount_to_precision(symbol, raw_amount))
    min_amt = None
    try:
        min_amt = market.get("limits", {}).get("amount", {}).get("min")
    except Exception:
        pass
    if min_amt is not None and amount < float(min_amt):
        raise RuntimeError(f"주문수량이 최소수량보다 작습니다. 계산수량={amount}, 최소수량={min_amt}")
    if amount <= 0:
        raise RuntimeError("주문수량이 0 이하입니다.")
    return amount

def price_to_precision(exchange, symbol, price):
    return float(exchange.price_to_precision(symbol, price))

def get_limit_buy_price(exchange, symbol):
    current_price = get_current_futures_price(exchange, symbol)
    return price_to_precision(exchange, symbol, current_price * BUY_LIMIT_PRICE_MULTIPLIER)

def fetch_order_status_safe(exchange, order_id, symbol):
    try:
        return exchange.fetch_order(order_id, symbol)
    except Exception as e:
        print(f"[주문조회 실패] {e}")
        return None

def wait_order_filled(exchange, order_id, symbol, max_check_count=ORDER_CHECK_COUNT, wait_sec=ORDER_WAIT_SEC):
    last_info = None
    for i in range(max_check_count):
        order_info = fetch_order_status_safe(exchange, order_id, symbol)
        last_info = order_info
        if order_info:
            status = str(order_info.get("status", "")).lower()
            filled = float(order_info.get("filled") or 0)
            amount = float(order_info.get("amount") or 0)
            print(f"[주문체결확인 {i+1}/{max_check_count}] status={status}, filled={filled}, amount={amount}")
            if status in ("closed", "filled"):
                return True, order_info
            if amount > 0 and filled >= amount:
                return True, order_info
        time.sleep(wait_sec)
    return False, last_info

def cancel_order_safe(exchange, order_id, symbol):
    try:
        result = exchange.cancel_order(order_id, symbol)
        print("[미체결 주문 취소]", result)
        return result
    except Exception as e:
        print(f"[주문취소 실패] {e}")
        return None

def place_open_long(exchange, symbol, amount):
    params = {"reduceOnly": False, "marginMode": MARGIN_MODE, "posSide": POSITION_SIDE}
    if USE_LIMIT_ORDER:
        price = get_limit_buy_price(exchange, symbol)
        order = exchange.create_limit_buy_order(symbol, amount, price, params=params)
        return order, price
    order = exchange.create_market_buy_order(symbol, amount, params=params)
    return order, None

def 체결가_조회(exchange, symbol, order_result, fallback_price, fallback_qty):
    if order_result:
        try:
            avg = order_result.get("average")
            filled = order_result.get("filled", 0) or 0
            cost = order_result.get("cost", 0) or 0
            if avg and float(avg) > 0 and float(filled) > 0:
                avg_f = float(avg); filled_f = float(filled)
                cost_f = float(cost) if cost else avg_f * filled_f
                return avg_f, filled_f, cost_f
        except Exception as e:
            print(f"[체결가 즉시 추출 실패] {e}")
    order_id = None
    try:
        if order_result:
            order_id = order_result.get("id")
    except Exception:
        pass
    if order_id:
        for attempt in range(1, FILL_FETCH_RETRIES + 1):
            try:
                time.sleep(FILL_FETCH_DELAY)
                fetched = exchange.fetch_order(order_id, symbol)
                if fetched:
                    avg = fetched.get("average")
                    filled = fetched.get("filled", 0) or 0
                    cost = fetched.get("cost", 0) or 0
                    if avg and float(avg) > 0 and float(filled) > 0:
                        avg_f = float(avg); filled_f = float(filled)
                        cost_f = float(cost) if cost else avg_f * filled_f
                        return avg_f, filled_f, cost_f
            except Exception as e:
                if attempt == FILL_FETCH_RETRIES:
                    print(f"[fetch_order 최종 실패 {attempt}/{FILL_FETCH_RETRIES}] {e}")
    print(f"⚠️ 체결가 조회 실패 → fallback 사용 (가격={fallback_price}, 수량={fallback_qty})")
    return float(fallback_price), float(fallback_qty), float(fallback_price) * float(fallback_qty)

# =========================
# ✅ 상태 관리 및 엑셀 함수
# =========================
def load_round_state():
    if not os.path.exists(ROUND_STATE_FILE):
        return {}
    try:
        with open(ROUND_STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"[상태파일 로드 실패] {e}")
        return {}

def save_round_state(state):
    try:
        d = os.path.dirname(ROUND_STATE_FILE)
        if d and not os.path.exists(d):
            os.makedirs(d, exist_ok=True)
        with open(ROUND_STATE_FILE, "w", encoding="utf-8") as f:
            json.dump(state, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[상태파일 저장 실패] {e}")

def update_round_stage(stock_code, round_number, stage, preservation_line=None, target_price=None, buy_price=None, quantity=None):
    state = load_round_state()
    rounds = state.setdefault(stock_code, {}).setdefault("rounds", {})
    r_key = str(int(round_number))
    r_info = rounds.setdefault(r_key, {})

    r_info["stage"] = int(stage)
    if preservation_line is not None:
        r_info["preservation_line"] = float(preservation_line)
    if target_price is not None:
        r_info["target_price"] = float(target_price)
    if buy_price is not None:
        r_info["buy_price"] = float(buy_price)
    if quantity is not None:
        r_info["quantity"] = float(quantity)
    if stage == 2 and not r_info.get("stage2_entered_at"):
        r_info["stage2_entered_at"] = datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S")

    save_round_state(state)
    return r_info

from openpyxl import load_workbook, Workbook

def 투자금시트_초기화_및_불러오기(stock_code, default_invest_money):
    default_invest_money = int(default_invest_money)
    if not os.path.exists(FILENAME):
        wb = Workbook()
        ws_trade = wb.active
        ws_trade.title = "매매기록"
        ws_trade.append(["차수", "날짜", "종목", "매수가", "수량", "고점대비하락조건", "매도가", "매도일", "매도수량"])
        ws = wb.create_sheet(INVEST_SHEET_NAME)
        ws.append(["종목", "투자금", "업데이트일", "마지막수익률(%)", "조정률(%)", "다음투자금"])
        ws.append([stock_code, int(default_invest_money), datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S"), "", "", int(default_invest_money)])
        wb.save(FILENAME)
        return int(default_invest_money)

    wb = load_workbook(FILENAME)
    if INVEST_SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(INVEST_SHEET_NAME)
        ws.append(["종목", "투자금", "업데이트일", "마지막수익률(%)", "조정률(%)", "다음투자금"])
        ws.append([stock_code, int(default_invest_money), datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S"), "", "", int(default_invest_money)])
        wb.save(FILENAME)
        return int(default_invest_money)

    ws = wb[INVEST_SHEET_NAME]
    stock_row = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == stock_code:
            stock_row = r
            break

    if stock_row is None:
        ws.append([stock_code, int(default_invest_money), datetime.now(KST).strftime("%Y-%m-%d %H:%M:%S"), "", "", int(default_invest_money)])
        wb.save(FILENAME)
        return int(default_invest_money)

    invest_val = ws.cell(row=stock_row, column=2).value
    try:
        invest_val = int(float(invest_val))
    except Exception:
        invest_val = int(default_invest_money)
    return int(invest_val)

def 매매기록불러오기():
    try:
        wb = load_workbook(FILENAME, read_only=True)
        sheet_name = "매매기록" if "매매기록" in wb.sheetnames else 0
        df = pd.read_excel(FILENAME, engine="openpyxl", sheet_name=sheet_name)
        if "매도일" in df.columns:
            df["매도일"] = df["매도일"].astype(str)
        if "날짜" in df.columns:
            df["날짜"] = pd.to_datetime(df["날짜"], errors="coerce").dt.strftime("%Y-%m-%d")
        if "고점대비하락조건" in df.columns:
            def _to_bool(v):
                if isinstance(v, bool): return v
                if v is None: return False
                s = str(v).strip().lower()
                return s in ("true", "t", "1", "y", "yes")
            df["고점대비하락조건"] = df["고점대비하락조건"].apply(_to_bool)
    except FileNotFoundError:
        df = pd.DataFrame(columns=["차수", "날짜", "종목", "매수가", "수량", "고점대비하락조건", "매도가", "매도일", "매도수량"])
    except Exception:
        df = pd.DataFrame(columns=["차수", "날짜", "종목", "매수가", "수량", "고점대비하락조건", "매도가", "매도일", "매도수량"])
    return df

def 매매기록저장(df):
    if "매도일" in df.columns:
        df["매도일"] = df["매도일"].astype(str)
    if os.path.exists(FILENAME):
        wb = load_workbook(FILENAME)
        if "매매기록" not in wb.sheetnames:
            wb.create_sheet("매매기록", 0)
            wb.save(FILENAME)
        with pd.ExcelWriter(FILENAME, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            df.to_excel(writer, index=False, sheet_name="매매기록")
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "매매기록"
        ws.append(["차수", "날짜", "종목", "매수가", "수량", "고점대비하락조건", "매도가", "매도일", "매도수량"])
        wb.save(FILENAME)
        with pd.ExcelWriter(FILENAME, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            df.to_excel(writer, index=False, sheet_name="매매기록")

def 과거_고점찾기(stock_code, ohlcv=None):
    if ohlcv is None:
        ohlcv = LsCommon.GetOhlcv("US", stock_code)
    if len(ohlcv) >= LOOKBACK_DAYS:
        recent_ohlcv = ohlcv[-LOOKBACK_DAYS:]
    else:
        recent_ohlcv = ohlcv
    return float(recent_ohlcv["close"].max())

def 가능다음차수(stock_data):
    for round_num in range(1, 40):
        if round_num not in stock_data["차수"].values or (
            round_num in stock_data["차수"].values and pd.notna(stock_data.loc[stock_data["차수"] == round_num, "매도가"]).all()
        ):
            return round_num
    return 1

def 매수기록(stock_code, round_number, buy_price, quantity, high_drop_condition):
    df = 매매기록불러오기()
    buy_date = datetime.now(KST).strftime("%Y-%m-%d")
    new_data = {
        "차수": round_number,
        "날짜": buy_date,
        "종목": stock_code,
        "매수가": buy_price,
        "수량": quantity,
        "고점대비하락조건": high_drop_condition,
        "매도가": None,
        "매도일": None,
        "매도수량": None
    }
    df = pd.concat([df, pd.DataFrame([new_data])], ignore_index=True)
    매매기록저장(df)

    target_price = float(buy_price) * (1.0 + SELL_PROFIT_PCT)
    update_round_stage(
        stock_code,
        round_number,
        stage=1,
        preservation_line=target_price,
        target_price=target_price,
        buy_price=float(buy_price),
        quantity=float(quantity)
    )
    print(f"✅ {stock_code} {round_number}차 매수 기록 및 상태(1단계) 등록 완료 (매수가: {buy_price}, 수량: {quantity})")

# =========================
# ✅ 시세 조회 폴백 함수
# =========================
# =========================
# ✅ 시세 데이터 재시도 및 폴백 래퍼 (bitget_Ejftkdhvkf.py 동일)
# =========================
def _값이_비정상(v):
    try:
        if v is None:
            return True
        if isinstance(v, str):
            s = v.strip().lower()
            if s in ("", "none", "nan", "nat"):
                return True
        if pd.isna(v):
            return True
        fv = float(v)
        if fv <= 0:
            return True
        return False
    except Exception:
        return True

def _최근_close_유효한지(ohlcv):
    try:
        if ohlcv is None or len(ohlcv) == 0:
            return False
        if "close" not in ohlcv.columns:
            return False
        recent_close = ohlcv["close"].tail(30)
        recent_close = pd.to_numeric(recent_close, errors="coerce").dropna()
        recent_close = recent_close[recent_close > 0]
        return len(recent_close) > 0
    except Exception:
        return False

def _장체크용_ohlcv_유효성체크(ohlcv):
    try:
        if ohlcv is None:
            return False, "OHLCV가 None"
        if not isinstance(ohlcv, pd.DataFrame):
            return False, f"OHLCV 타입 이상: {type(ohlcv)}"
        if len(ohlcv) == 0:
            return False, "OHLCV가 비어 있음"
        if "open" not in ohlcv.columns:
            return False, "OHLCV에 open 컬럼이 없음"
        if "close" not in ohlcv.columns:
            return False, "OHLCV에 close 컬럼이 없음"
        last_open = pd.to_numeric(pd.Series([ohlcv["open"].iloc[-1]]), errors="coerce").iloc[0]
        last_close = pd.to_numeric(pd.Series([ohlcv["close"].iloc[-1]]), errors="coerce").iloc[0]
        if _값이_비정상(last_open):
            return False, f"시가 비정상값: {last_open}"
        if _값이_비정상(last_close):
            return False, f"종가 비정상값: {last_close}"
        return True, "OK"
    except Exception as e:
        return False, str(e)

_원본_GetCurrentPrice = LsiUS.GetCurrentPrice
_원본_GetOhlcv = LsiUS.GetOhlcv

def _get_yf_current_price(stock_code):
    try:
        import requests
        url = f"https://query1.finance.yahoo.com/v8/finance/chart/{stock_code}"
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        res = requests.get(url, headers=headers, params={"range": "1d", "interval": "1m"}, timeout=10)
        if res.status_code == 200:
            result = res.json().get("chart", {}).get("result", [])
            if result:
                meta = result[0].get("meta", {})
                price = meta.get("regularMarketPrice") or meta.get("chartPreviousClose") or meta.get("previousClose")
                if price and not _값이_비정상(price):
                    return float(price)
                quotes = result[0].get("indicators", {}).get("quote", [{}])[0].get("close", [])
                valid_closes = [c for c in quotes if c is not None and not _값이_비정상(c)]
                if valid_closes:
                    return float(valid_closes[-1])
    except Exception as e:
        print(f"[야후 현재가 폴백 조회 오류] {e}")
    return None

def _GetCurrentPrice_with_retry(stock_code, *args, **kwargs):
    last_err = None
    for attempt in range(1, DATA_RETRY_COUNT + 1):
        try:
            current_price = _원본_GetCurrentPrice(stock_code, *args, **kwargs)
            if not _값이_비정상(current_price):
                return float(current_price)
            last_err = f"현재가 비정상값: {current_price}"
        except Exception as e:
            last_err = str(e)
        print(f"[LS 현재가 재시도 {attempt}/{DATA_RETRY_COUNT}] {stock_code} - {last_err}")
        time.sleep(DATA_RETRY_SLEEP_SEC)

    yf_price = _get_yf_current_price(stock_code)
    if yf_price is not None:
        print(f"[시세소스] {stock_code} 현재가 ← 야후 직접HTTP 폴백: {yf_price}")
        return float(yf_price)

    raise RuntimeError(f"{stock_code} LS 현재가 조회 실패: {last_err}")

def _GetOhlcv_with_retry(*args, **kwargs):
    market = args[0] if len(args) > 0 else kwargs.get("market", "UNKNOWN")
    stock_code = args[1] if len(args) > 1 else kwargs.get("stock_code", "UNKNOWN")
    limit = args[2] if len(args) > 2 else kwargs.get("limit", None)
    last_err = None

    for attempt in range(1, DATA_RETRY_COUNT + 1):
        try:
            _kw = dict(kwargs)
            if not (limit is not None and int(limit) <= 10):
                _kw.setdefault("source", "yf_auto")
            ohlcv = _원본_GetOhlcv(*args, **_kw)

            if ohlcv is None:
                last_err = "OHLCV가 None"
            elif not isinstance(ohlcv, pd.DataFrame):
                last_err = f"OHLCV 타입 이상: {type(ohlcv)}"
            elif len(ohlcv) == 0:
                last_err = "OHLCV가 비어 있음"
            elif "open" not in ohlcv.columns:
                last_err = "OHLCV에 open 컬럼이 없음"
            elif "close" not in ohlcv.columns:
                last_err = "OHLCV에 close 컬럼이 없음"
            else:
                last_open = ohlcv["open"].iloc[-1]
                if _값이_비정상(last_open):
                    last_err = f"시가 비정상값: {last_open}"
                elif not _최근_close_유효한지(ohlcv):
                    last_err = "최근 30개 close 데이터가 유효하지 않음"
                else:
                    return ohlcv
        except Exception as e:
            last_err = str(e)

        print(f"[OHLCV 재시도 {attempt}/{DATA_RETRY_COUNT}] {market} {stock_code} - {last_err}")
        time.sleep(DATA_RETRY_SLEEP_SEC)

    raise RuntimeError(f"{stock_code} OHLCV 조회 실패: {last_err}")

# LS 데이터 함수를 재시도/폴백 버전으로 패치
LsiUS.GetCurrentPrice = _GetCurrentPrice_with_retry
LsCommon.GetOhlcv = _GetOhlcv_with_retry

def get_market_prices(stock_code, exchange, symbol):
    """LS 시세 조회 (실패 시 야후 및 비트겟 선물 시세 폴백)"""
    ohlcv = LsCommon.GetOhlcv("US", stock_code)
    current_price = None
    try:
        current_price = LsiUS.GetCurrentPrice(stock_code)
    except Exception as e:
        print(f"[현재가 조회 실패] {e}")

    futures_price = get_current_futures_price(exchange, symbol)
    if current_price is None:
        current_price = futures_price
        print(f"[시세소스] 비트겟 선물 현재가 대체 사용: {current_price}")

    open_price = float(ohlcv["open"].iloc[-1])
    return current_price, open_price, futures_price, ohlcv

# =========================
# ✅ 수동 매수 메인 프로세스
# =========================
def execute_manual_buy(check_only: bool = False, force: bool = False):
    print("=" * 60)
    print("🚀 [비트겟 떨사오팔] 수동 매수 점검 및 실행 스크립트 시작")
    print(f"⏰ 실행시각: {datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S')} KST")
    print(f"⚙️ 실행모드: {'조회 전용 (--check-only)' if check_only else ('강제 매수 (--force)' if force else '기준 충족 시 매수 (기본)')}")
    print("=" * 60)

    LsCommon.SetChangeMode("REAL")
    exchange = create_exchange()
    symbol = resolve_symbol(exchange)
    stock_code = "SOXL"

    # 1. 엑셀 데이터 및 다음 차수 확인
    df = 매매기록불러오기()
    stock_data = df[df["종목"] == stock_code] if not df.empty else pd.DataFrame()
    next_round = 가능다음차수(stock_data) if not stock_data.empty else 1
    base_invest_money = 투자금시트_초기화_및_불러오기(stock_code, INITIAL_INVEST_USDT)

    # 2. 시세 데이터 조회
    current_price, open_price, futures_price, ohlcv = get_market_prices(stock_code, exchange, symbol)
    highest_price = 과거_고점찾기(stock_code, ohlcv)

    # 3. 매수 조건 계산
    target_buy_threshold = open_price * (1.0 - BUY_DROP_PCT)
    drop_pct_from_open = ((current_price - open_price) / open_price) * 100.0
    high_drop_condition = ((highest_price - current_price) / highest_price >= HIGH_DROP_THRESHOLD) if highest_price else False
    high_drop_rate = ((highest_price - current_price) / highest_price * 100.0) if highest_price else 0.0

    # 4. 투자금 및 가용 잔고 계산
    invest_multiplier = HIGH_DROP_INVEST_MULTIPLIER if high_drop_condition else 1.0
    adjusted_investment = base_invest_money * invest_multiplier
    available_usdt = get_available_usdt(exchange)

    # 매수 조건 충족 여부
    is_condition_met = (current_price <= target_buy_threshold)

    # 상태 브리핑 출력
    print(f"\n📊 [시세 및 조건 분석 보고]")
    print(f"  • 대상 종목: {stock_code} (선물 심볼: {symbol})")
    print(f"  • 다음 진입 차수: {next_round}차")
    print(f"  • 당일 시가(Open): {open_price:.4f} USD")
    print(f"  • 현재가(Price)  : {current_price:.4f} USD (시가대비: {drop_pct_from_open:+.2f}%)")
    print(f"  • 매수 기준가    : {target_buy_threshold:.4f} USD (시가 -{BUY_DROP_PCT*100:.1f}%)")
    print(f"  • 40일 최고가    : {highest_price:.4f} USD (낙폭: {high_drop_rate:.2f}%)")
    print(f"  • 대폭락 증액조건: {'충족 (1.5배 증액)' if high_drop_condition else '미충족 (기본금액)'}")
    print(f"  • 필요 투자금    : 약 {adjusted_investment:.1f} USDT (기본 {base_invest_money} USDT)")
    print(f"  • 비트겟 가용잔고: {available_usdt:.2f} USDT")
    print(f"  • 매수 조건 부합 : {'✅ 조건 충족' if is_condition_met else '❌ 조건 미충족'}")

    # [조회 전용 모드]
    if check_only:
        print("\n🔎 [--check-only 모드] 주문을 실행하지 않고 상태 조회를 완료했습니다.")
        return

    # [매수 조건 미충족 시 처리]
    if not is_condition_met and not force:
        msg = (
            f"⏸️ [수동 매수 보류] {stock_code} 매수 조건 미충족\n"
            f"🔹 현재 진입 예정 차수: {next_round}차\n"
            f"📈 당일 시가: {open_price:.2f} | 💲 현재가: {current_price:.2f}\n"
            f"🎯 매수 기준가: {target_buy_threshold:.2f} (시가 대비 -{BUY_DROP_PCT*100:.1f}%)\n"
            f"📊 현재 변동률: {drop_pct_from_open:+.2f}%\n"
            f"💡 강제로 매수하시려면 `python manual_buy.py --force` 를 사용하세요."
        )
        print(f"\n⚠️ 매수 조건 미충족: 현재가({current_price:.4f}) > 기준가({target_buy_threshold:.4f})")
        print("💡 매수를 강제 집행하려면 '--force' 옵션을 붙여 실행하세요.")
        telegram_alert.SendMessage(msg)
        return

    # [가용 잔고 부족 점검]
    if 0 <= available_usdt < adjusted_investment:
        msg = (
            f"🚨 [보유금 부족] 떨사오팔 {stock_code} {next_round}차 수동 매수 보류\n"
            f"🔹 사유: 비트겟 선물 가용 USDT가 필요 투자금보다 부족합니다.\n"
            f"💵 필요 투자금: 약 {adjusted_investment:.1f} USDT\n"
            f"💰 가용 잔고: {available_usdt:.2f} USDT\n"
            f"💡 조치: 비트겟 선물 계좌에 USDT를 보충 후 다시 실행해 주세요."
        )
        print(f"\n❌ [보유금 부족] 필요: {adjusted_investment:.1f} USDT > 가용: {available_usdt:.2f} USDT")
        telegram_alert.SendMessage(msg)
        return

    # 5. 주문 수량 산출 및 선물 롱 매수 주문 실행
    print(f"\n⚡ [{ '강제' if force and not is_condition_met else '정상' } 수동 매수 실행 시작]")
    amount = usdt_to_amount(exchange, symbol, adjusted_investment, futures_price)
    print(f"📦 주문 수량: {amount} (계산 선물현재가: {futures_price:.4f})")

    try:
        order, order_price = place_open_long(exchange, symbol, amount)
        print(f"📨 주문 접수 완료:", order)
    except ccxt.InsufficientFunds as e:
        msg = (
            f"🚨 [보유금 부족] 떨사오팔 {stock_code} {next_round}차 수동 매수 실패\n"
            f"🔹 거래소 거절: 증거금 부족 (InsufficientFunds)\n"
            f"💵 주문 금액: 약 {adjusted_investment:.1f} USDT (수량: {amount})"
        )
        telegram_alert.SendMessage(msg)
        print(f"[주문 거절 - 잔고 부족] {e}")
        return
    except Exception as e:
        print(f"❌ [주문 실패] {e}")
        telegram_alert.SendMessage(f"🚨 [수동 매수 에러] {stock_code} {next_round}차 주문 실패: {e}")
        return

    # 6. 체결 확인 및 미체결 처리
    fallback_buy = float(order_price) if order_price is not None else float(futures_price)
    order_id = order.get("id") if isinstance(order, dict) else None
    if order_id:
        filled_ok, _ = wait_order_filled(exchange, order_id, symbol)
        if not filled_ok and USE_LIMIT_ORDER:
            cancel_order_safe(exchange, order_id, symbol)
            telegram_alert.SendMessage(f"⚠️ 떨사오팔(Bitget) {stock_code} {next_round}차 수동 매수 지정가 미체결 → 취소")
            print("❌ 지정가 주문 미체결로 취소되었습니다.")
            return

    # 7. 실제 체결가 및 수량 회수
    real_buy_price, real_buy_qty, _ = 체결가_조회(exchange, symbol, order, fallback_buy, amount)
    buy_amt = real_buy_qty

    # 8. 엑셀 매매기록 및 JSON 상태 동기화 저장
    매수기록(stock_code, next_round, real_buy_price, buy_amt, high_drop_condition)

    # 9. 텔레그램 및 구글 시트 알림
    invest_suffix = f" → 증액({HIGH_DROP_INVEST_MULTIPLIER:.1f}배) 적용" if high_drop_condition else ""
    target_sell_price = real_buy_price * (1.0 + SELL_PROFIT_PCT)

    msg = (
        f"🟢 [수동 매수 실행] 떨사오팔 {stock_code} {next_round}차 매수 완료\n"
        f"🔹 구분: {'수동 강제 매수 (--force)' if force and not is_condition_met else '수동 조건 매수'}\n"
        f"📈 당일 시가(LS): {open_price:.2f} | 💲 체결시 현재가: {current_price:.2f}\n"
        f"📉 매수체결가(Bitget): {real_buy_price:.4f}\n"
        f"🎯 1차 목표가(+{SELL_PROFIT_PCT*100:.1f}%): {target_sell_price:.4f}\n"
        f"💵 투자금: {base_invest_money}{invest_suffix} USDT\n"
        f"📊 체결 수량: {buy_amt}주\n"
        f"📅 매수일: {datetime.now(KST).strftime('%Y-%m-%d')}\n"
        f"📊 고점 대비 낙폭률: {high_drop_rate:.2f}%\n"
        f"💡 24시간 봇이 자동으로 해당 차수의 매도 감시(목표가 도달 시 익절)를 시작합니다."
    )
    telegram_alert.SendMessage(msg)

    # 구글 시트 웹훅 전송
    try:
        sheet_logger.log_trade(
            action="BUY",
            symbol=symbol,
            level=f"{next_round}차(수동)",
            price=real_buy_price,
            amount=buy_amt,
            note=f"떨사 수동 매수 실행 (낙폭률 {high_drop_rate:.2f}%)"
        )
    except Exception as e:
        print(f"[구글 시트 로깅 경고] {e}")

    print("\n🎉 [수동 매수 완료] 모든 처리가 정상적으로 완료되었습니다.")
    print(f"👉 엑셀 매매기록 및 {ROUND_STATE_FILE} 에 {next_round}차가 등록되었습니다.")
    print("👉 24시간 실행 중인 bitget_Ejftkdhvkf.py 가 다음 턴부터 이 포지션을 자동 매도 감시합니다.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="비트겟 떨사오팔 수동 매수 전용 스크립트")
    parser.add_argument("--check-only", action="store_true", help="주문 없이 현재 매수 조건 및 가용 잔고만 조회")
    parser.add_argument("--force", action="store_true", help="매수 조건(시가 -1.5%% 하락) 미달이어도 즉시 1회 강제 매수")
    args = parser.parse_args()

    execute_manual_buy(check_only=args.check_only, force=args.force)
